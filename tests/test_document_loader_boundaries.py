"""Offline file-format, parser-failure, OCR, PDF-page, and ZIP cleanup checks."""

import builtins
import sys
import zipfile
from pathlib import Path
from types import ModuleType, SimpleNamespace

import pytest


def test_empty_reason_labels_cover_archives_pdfs_images_and_plain_text():
    import rag.document_loader as loader

    assert loader._empty_reason(".zip")[0] == loader.REASON_EMPTY_CONTENT
    assert loader._empty_reason(".pdf")[0] == loader.REASON_EMPTY_CONTENT
    assert loader._empty_reason(".docx")[0] == loader.REASON_EMPTY_CONTENT
    assert loader._empty_reason(".png")[0] == loader.REASON_OCR_NO_TEXT
    assert loader._empty_reason(".txt")[0] == loader.REASON_EMPTY_CONTENT


def test_text_loader_tries_legacy_encodings_and_reports_missing_file(tmp_path):
    import rag.document_loader as loader

    gbk = tmp_path / "legacy.txt"
    gbk.write_bytes("你好".encode("gbk"))
    assert loader.load_txt_md(str(gbk)) == "你好"

    undecodable_utf8 = tmp_path / "latin.txt"
    undecodable_utf8.write_bytes(b"\xff")
    assert loader.load_txt_md(str(undecodable_utf8)) == "\xff"

    with pytest.raises(FileNotFoundError):
        loader.load_txt_md(str(tmp_path / "missing.txt"))


def test_docx_and_xlsx_parsers_extract_text_and_report_corrupt_inputs(tmp_path):
    from docx import Document
    from openpyxl import Workbook

    import rag.document_loader as loader

    docx_path = tmp_path / "notes.docx"
    doc = Document()
    doc.add_paragraph("first paragraph")
    doc.add_paragraph("")
    doc.add_paragraph("second paragraph")
    doc.save(docx_path)
    assert loader.load_docx(str(docx_path)) == "first paragraph\n\nsecond paragraph"

    xlsx_path = tmp_path / "table.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["name", None, "count"])
    sheet.append(["item", 3, None])
    workbook.create_sheet("Empty")
    workbook.save(xlsx_path)
    text = loader.load_xlsx(str(xlsx_path))
    assert "--- Sheet: Data ---\nname | count\nitem | 3" in text
    assert "--- Sheet: Empty ---" in text

    broken = tmp_path / "broken.docx"
    broken.write_text("not a word document", encoding="utf-8")
    outcome = loader.load_document_detailed(str(broken))
    assert outcome.status == "failed"
    assert outcome.reason == loader.REASON_PARSE_ERROR


def test_xlsx_archive_is_closed_when_sheet_reading_fails(monkeypatch, tmp_path):
    import openpyxl

    import rag.document_loader as loader

    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"fake")
    closed = []

    class Sheet:
        def iter_rows(self, *, values_only):
            assert values_only is True
            raise RuntimeError("sheet read failed")

    class Workbook:
        def __init__(self):
            self.sheetnames = ["Broken"]

        def __getitem__(self, _name):
            return Sheet()

        def close(self):
            closed.append(True)

    monkeypatch.setattr(openpyxl, "load_workbook", lambda *_args, **_kwargs: Workbook())
    outcome = loader.load_document_detailed(str(path))
    assert outcome.status == "failed"
    assert outcome.reason == loader.REASON_PARSE_ERROR
    assert closed == [True]


def test_pdf_pages_preserve_blank_page_positions_and_detailed_result(monkeypatch, tmp_path):
    import rag.document_loader as loader

    path = tmp_path / "pages.pdf"
    path.write_bytes(b"fake pdf")

    class Page:
        def __init__(self, text):
            self.text = text

        def extract_text(self):
            return self.text

    monkeypatch.setitem(
        sys.modules,
        "pypdf",
        SimpleNamespace(PdfReader=lambda _path: SimpleNamespace(pages=[Page("first"), Page(None), Page("third")])),
    )
    assert loader.load_pdf_pages(str(path)) == ["first", "", "third"]
    assert loader.load_pdf(str(path)) == "first\n\nthird"
    outcome = loader.load_document_detailed(str(path))
    assert outcome.status == "ok"
    assert outcome.content == "first\n\nthird"
    assert outcome.pages == ["first", "", "third"]


def test_pdf_empty_content_and_parser_errors_are_classified(monkeypatch, tmp_path):
    import rag.document_loader as loader

    path = tmp_path / "empty.pdf"
    path.write_bytes(b"fake")
    monkeypatch.setitem(
        sys.modules,
        "pypdf",
        SimpleNamespace(PdfReader=lambda _path: SimpleNamespace(pages=[])),
    )
    outcome = loader.load_document_detailed(str(path))
    assert outcome.status == "empty"
    assert outcome.reason == loader.REASON_EMPTY_CONTENT

    def broken_reader(_path):
        raise ValueError("malformed PDF")

    monkeypatch.setitem(sys.modules, "pypdf", SimpleNamespace(PdfReader=broken_reader))
    outcome = loader.load_document_detailed(str(path))
    assert outcome.status == "failed"
    assert outcome.reason == loader.REASON_PARSE_ERROR
    assert "malformed PDF" in outcome.detail


def test_unsupported_formats_keep_compatible_exception_and_detailed_reason(tmp_path):
    import rag.document_loader as loader

    path = tmp_path / "notes.unsupported"
    path.write_text("data", encoding="utf-8")
    with pytest.raises(loader.UnsupportedFormatError):
        loader.load_single_document(str(path))
    outcome = loader.load_document_detailed(str(path))
    assert outcome.status == "failed"
    assert outcome.reason == loader.REASON_UNSUPPORTED_FORMAT


def test_missing_ocr_dependencies_are_reported_as_engine_failure(monkeypatch, tmp_path):
    import rag.document_loader as loader

    real_import = builtins.__import__

    def import_without_ocr(name, *args, **kwargs):
        if name in {"pytesseract", "PIL"}:
            raise ImportError("OCR package missing")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", import_without_ocr)
    outcome = loader.load_document_detailed(str(tmp_path / "photo.png"))
    assert outcome.status == "failed"
    assert outcome.reason == loader.REASON_FILE_MISSING

    image = tmp_path / "photo.png"
    image.write_bytes(b"image")
    outcome = loader.load_document_detailed(str(image))
    assert outcome.reason == loader.REASON_OCR_ENGINE_MISSING


def test_ocr_engine_lookup_and_corrupt_image_failures_are_distinguished(monkeypatch, tmp_path):
    import rag.document_loader as loader

    not_found = type("TesseractNotFoundError", (RuntimeError,), {})

    class ImageContext:
        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return None

    def image_open(_path):
        return ImageContext()

    def ocr_missing(*_args, **_kwargs):
        raise not_found("missing executable")

    tesseract = ModuleType("pytesseract")
    tesseract.TesseractNotFoundError = not_found
    tesseract.image_to_string = ocr_missing
    pil = ModuleType("PIL")
    pil.Image = SimpleNamespace(open=image_open)
    monkeypatch.setitem(sys.modules, "pytesseract", tesseract)
    monkeypatch.setitem(sys.modules, "PIL", pil)

    with pytest.raises(loader.MissingOcrEngineError):
        loader.load_image_ocr(str(tmp_path / "photo.png"))

    tesseract.image_to_string = lambda *_args, **_kwargs: (_ for _ in ()).throw(
        ValueError("corrupt image")
    )
    with pytest.raises(loader.DocumentLoadError, match="corrupt image"):
        loader.load_image_ocr(str(tmp_path / "photo.png"))


def test_zip_extracts_supported_members_and_always_cleans_temporary_directory(
    monkeypatch, tmp_path
):
    import shutil

    import rag.document_loader as loader

    archive = tmp_path / "docs.zip"
    with zipfile.ZipFile(archive, "w") as zf:
        zf.writestr("nested/readme.txt", "hello")
    extracted = tmp_path / "extract"
    removed = []
    original_rmtree = shutil.rmtree

    def tracking_rmtree(path, ignore_errors=False):
        removed.append(Path(path))
        original_rmtree(path, ignore_errors=ignore_errors)

    monkeypatch.setattr(loader.tempfile, "mkdtemp", lambda **_kwargs: str(extracted))
    monkeypatch.setattr(shutil, "rmtree", tracking_rmtree)
    monkeypatch.setattr(loader, "load_single_document", lambda path, _upload: Path(path).read_text())

    assert loader.load_zip(str(archive), str(tmp_path)) == ["hello"]
    assert removed == [extracted]
    assert not extracted.exists()


def test_empty_and_corrupt_zip_outcomes_are_distinguishable(tmp_path):
    import rag.document_loader as loader

    empty_zip = tmp_path / "empty.zip"
    with zipfile.ZipFile(empty_zip, "w"):
        pass
    outcome = loader.load_document_detailed(str(empty_zip))
    assert outcome.status == "empty"
    assert outcome.reason == loader.REASON_EMPTY_CONTENT

    broken_zip = tmp_path / "broken.zip"
    broken_zip.write_bytes(b"not zip")
    outcome = loader.load_document_detailed(str(broken_zip))
    assert outcome.status == "failed"
    assert outcome.reason == loader.REASON_PARSE_ERROR
